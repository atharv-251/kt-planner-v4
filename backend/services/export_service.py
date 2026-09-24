from pathlib import Path
from typing import Dict, Any, List
import io
import csv
from datetime import date, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from backend.models.transition import Transition, ProjectProfile
from backend.models.knowledge import KnowledgeNode, KTLevelEvaluation
from backend.models.scheduling import KTSession
from backend.models.stakeholder import Stakeholder
from backend.models.governance import PlanPatch
from backend.services.capacity_service import CapacityService
from backend.services.validation_service import ValidationService

class ExportService:
    @staticmethod
    def generate_excel_master_package(db: Session, transition_id: str) -> io.BytesIO:
        transition = db.query(Transition).filter(Transition.id == transition_id).first()
        if not transition:
            raise ValueError(f"Transition {transition_id} not found")

        wb = Workbook()
        # Remove default sheet
        default_sheet = wb.active

        # Style helpers
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Deep Blue
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
        sub_font = Font(name="Segoe UI", size=10, italic=True, color="555555")
        data_font = Font(name="Segoe UI", size=10)
        thin_border = Border(
            left=Side(style="thin", color="E5E7EB"),
            right=Side(style="thin", color="E5E7EB"),
            top=Side(style="thin", color="E5E7EB"),
            bottom=Side(style="thin", color="E5E7EB"),
        )

        def style_headers(ws, headers: List[str], start_row: int = 3):
            for col_num, h in enumerate(headers, 1):
                cell = ws.cell(row=start_row, column=col_num, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[start_row].height = 25

        def auto_fit_cols(ws):
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val = str(cell.value or "")
                    if len(val) > max_len:
                        max_len = len(val)
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 60)

        # -------------------------------------------------------------
        # Sheet 1: Project Profile & Executive Summary
        # -------------------------------------------------------------
        ws_summary = wb.create_sheet(title="Executive Summary")
        ws_summary.views.sheetView[0].showGridLines = True
        ws_summary.cell(row=1, column=1, value="KT PLANNER - EXECUTIVE TRANSITION SUMMARY").font = title_font
        ws_summary.cell(row=2, column=1, value=f"Generated on {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}").font = sub_font

        profile = db.query(ProjectProfile).filter(ProjectProfile.transition_id == transition_id).first()
        cap = CapacityService.evaluate_capacity_balance(db, transition_id)

        rows_summary = [
            ("Transition Name", transition.name),
            ("Application Category", transition.category),
            ("Primary Country & Timezone", f"{transition.primary_country} ({transition.timezone})"),
            ("Start & End Dates", f"{transition.start_date} to {transition.end_date}"),
            ("Transition Duration (Total Days)", transition.total_duration_days),
            ("Shadow / Reverse Shadow Days", f"{transition.shadow_days} / {transition.reverse_shadow_days} days"),
            ("Available KT Days", cap["available_kt_days"]),
            ("Target KT Capacity", f"{cap['target_capacity_hours']} Hours"),
            ("Generated Topic Capacity", f"{cap['generated_hours']} Hours ({cap['balance_ratio_percent']}%)"),
            ("Capacity Balance Status", cap["status"]),
            ("Business Criticality", profile.criticality if profile else "Business Critical"),
            ("Support Model", profile.support_model if profile else "AMS 2"),
            ("Approval Status", "Approved" if (profile and profile.is_approved) else "Draft Review"),
        ]
        style_headers(ws_summary, ["Parameter", "Transition Value"], start_row=4)
        for idx, (k, v) in enumerate(rows_summary, start=5):
            c1 = ws_summary.cell(row=idx, column=1, value=k)
            c2 = ws_summary.cell(row=idx, column=2, value=str(v))
            c1.font = Font(name="Segoe UI", size=10, bold=True)
            c2.font = data_font
            c1.border = thin_border
            c2.border = thin_border

        # -------------------------------------------------------------
        # Sheet 2: KT Master Hierarchy
        # -------------------------------------------------------------
        ws_hierarchy = wb.create_sheet(title="KT Master Plan")
        ws_hierarchy.views.sheetView[0].showGridLines = True
        ws_hierarchy.cell(row=1, column=1, value="KNOWLEDGE TRANSFER MASTER HIERARCHY").font = title_font
        
        headers_hier = [
            "Level Type", "Topic / Subtopic Name", "Category", "Est. Hours",
            "Delivery Method", "Weightage %", "Evidence Sources", "Description"
        ]
        style_headers(ws_hierarchy, headers_hier, start_row=3)

        nodes = db.query(KnowledgeNode).filter(KnowledgeNode.transition_id == transition_id).order_by(KnowledgeNode.order_index).all()
        for idx, n in enumerate(nodes, start=4):
            ev_str = "; ".join(n.evidence_references or [])
            row_vals = [
                n.node_type.upper(), n.name, n.category, n.estimated_hours,
                n.recommended_method, f"{n.weightage_percent}%", ev_str, n.description
            ]
            for c_idx, val in enumerate(row_vals, 1):
                cell = ws_hierarchy.cell(row=idx, column=c_idx, value=val)
                cell.font = data_font
                cell.border = thin_border

        # -------------------------------------------------------------
        # Sheet 3: KT Schedule
        # -------------------------------------------------------------
        ws_schedule = wb.create_sheet(title="KT Schedule")
        ws_schedule.views.sheetView[0].showGridLines = True
        ws_schedule.cell(row=1, column=1, value="CONFIRMABLE SESSION SCHEDULE").font = title_font

        headers_sched = [
            "Session Title", "KT Level", "Date", "Day", "Start Time",
            "End Time", "Duration (Hrs)", "Delivery Mode", "SME Assigned", "Receiver Assigned",
            "Required Attendees", "Optional Attendees", "Status", "Conflicts"
        ]
        style_headers(ws_schedule, headers_sched, start_row=3)

        sessions = (
            db.query(KTSession)
            .filter(KTSession.transition_id == transition_id)
            .order_by(KTSession.scheduled_date, KTSession.start_time)
            .all()
        )
        stakeholders = db.query(Stakeholder).filter(Stakeholder.transition_id == transition_id).all()
        stk_map = {s.id: s for s in stakeholders}
        level_weights = {"L1": 1, "L2": 2, "L3": 3}

        for idx, s in enumerate(sessions, start=4):
            sme_obj = stk_map.get(s.sme_id)
            rcv_obj = stk_map.get(s.receiver_id)
            sme_name = sme_obj.name if sme_obj else "Unassigned"
            rcv_name = rcv_obj.name if rcv_obj else "Unassigned"
            
            t_rank = level_weights.get(s.level, 1)
            # Same level SME and Team member emails -> Required
            req_emails = [stk.email for stk in stakeholders if stk.email and level_weights.get(stk.level, 1) == t_rank]
            # Higher SME and Team member emails -> Optional
            opt_emails = [stk.email for stk in stakeholders if stk.email and level_weights.get(stk.level, 1) > t_rank]

            conflict_str = "; ".join(s.conflict_flags or [])
            row_vals = [
                s.session_title, s.level, s.scheduled_date.isoformat(),
                s.scheduled_date.strftime("%A"), s.start_time.strftime("%H:%M"),
                s.end_time.strftime("%H:%M"), s.duration_hours, s.delivery_mode,
                sme_name, rcv_name,
                "; ".join(req_emails), "; ".join(opt_emails),
                s.status.upper(), conflict_str or "None"
            ]
            for c_idx, val in enumerate(row_vals, 1):
                cell = ws_schedule.cell(row=idx, column=c_idx, value=val)
                cell.font = data_font
                cell.border = thin_border

        # -------------------------------------------------------------
        # Sheet 4: KT Level Matrix (L1/L2/L3)
        # -------------------------------------------------------------
        ws_matrix = wb.create_sheet(title="Level Coverage Matrix")
        ws_matrix.views.sheetView[0].showGridLines = True
        ws_matrix.cell(row=1, column=1, value="L1/L2/L3 LEVEL EVALUATION MATRIX").font = title_font

        headers_matrix = [
            "Topic Name", "Level Scope", "Applicability", "Learning Objective",
            "Expected Outcome", "Evidence Justification"
        ]
        style_headers(ws_matrix, headers_matrix, start_row=3)

        evals = db.query(KTLevelEvaluation).filter(KTLevelEvaluation.transition_id == transition_id).all()
        node_map = {n.id: n.name for n in nodes}

        for idx, ev in enumerate(evals, start=4):
            row_vals = [
                node_map.get(ev.node_id, "Topic"), ev.level_scope, ev.applicability.upper(),
                ev.learning_objective, ev.expected_outcome, ev.evidence or ev.justification
            ]
            for c_idx, val in enumerate(row_vals, 1):
                cell = ws_matrix.cell(row=idx, column=c_idx, value=val)
                cell.font = data_font
                cell.border = thin_border

        # -------------------------------------------------------------
        # Sheet 5: Governance & Validation Scorecard
        # -------------------------------------------------------------
        ws_val = wb.create_sheet(title="Validation Report")
        ws_val.views.sheetView[0].showGridLines = True
        ws_val.cell(row=1, column=1, value="GOVERNANCE & READINESS VALIDATION REPORT").font = title_font

        val_report = ValidationService.run_full_validation(db, transition_id)
        ws_val.cell(row=2, column=1, value=f"Overall Status: {val_report.overall_status} | Score: {val_report.score_percent}%").font = sub_font

        headers_val = ["Governance Check", "Category", "Result", "Severity", "Compliance Message"]
        style_headers(ws_val, headers_val, start_row=4)

        for idx, chk in enumerate(val_report.checks, start=5):
            row_vals = [
                chk.check_name, chk.category.upper(), "PASS" if chk.passed else "FAIL",
                chk.severity.upper(), chk.message
            ]
            for c_idx, val in enumerate(row_vals, 1):
                cell = ws_val.cell(row=idx, column=c_idx, value=val)
                cell.font = data_font
                cell.border = thin_border

        # -------------------------------------------------------------
        # Sheet 6: Source Traceability Matrix
        # -------------------------------------------------------------
        ws_trace = wb.create_sheet(title="Traceability Matrix")
        ws_trace.views.sheetView[0].showGridLines = True
        ws_trace.cell(row=1, column=1, value="SOURCE EVIDENCE TRACEABILITY MATRIX").font = title_font

        headers_trace = ["Node Type", "Topic / Subtopic Name", "Category", "Source Document / Sheet Citation", "Evidence Verification"]
        style_headers(ws_trace, headers_trace, start_row=3)

        for idx, n in enumerate(nodes, start=4):
            sources = "; ".join(n.evidence_references or ["Transition Intake Workbook"])
            row_vals = [
                n.node_type.upper(),
                n.name,
                n.category,
                sources,
                "Verified Evidence" if n.evidence_references else "Synthesized Domain Knowledge",
            ]
            for c_idx, val in enumerate(row_vals, 1):
                cell = ws_trace.cell(row=idx, column=c_idx, value=val)
                cell.font = data_font
                cell.border = thin_border

        # -------------------------------------------------------------
        # Sheet 7: Calendar Availability Report
        # -------------------------------------------------------------
        ws_avail = wb.create_sheet(title="Availability Report")
        ws_avail.views.sheetView[0].showGridLines = True
        ws_avail.cell(row=1, column=1, value="STAKEHOLDER & HOLIDAY CALENDAR AVAILABILITY").font = title_font

        headers_avail = ["Date", "Day of Week", "Calendar Status", "Bank Holiday / Leave Reason", "Assigned Sessions"]
        style_headers(ws_avail, headers_avail, start_row=3)

        from backend.services.availability_service import AvailabilityService
        avail_data = AvailabilityService.get_transition_availability_matrix(db, transition_id)
        session_dates = {}
        for s in sessions:
            dt_key = s.scheduled_date.isoformat()
            session_dates[dt_key] = session_dates.get(dt_key, 0) + 1

        for idx, day in enumerate(avail_data.get("days", []), start=4):
            row_vals = [
                day["date"],
                day["day_of_week"],
                day["status"],
                day.get("holiday_name") or "Standard Working Day",
                f"{session_dates.get(day['date'], 0)} sessions booked",
            ]
            for c_idx, val in enumerate(row_vals, 1):
                cell = ws_avail.cell(row=idx, column=c_idx, value=val)
                cell.font = data_font
                cell.border = thin_border

        # -------------------------------------------------------------
        # Sheet 8: Version Audit Report
        # -------------------------------------------------------------
        ws_audit = wb.create_sheet(title="Version Audit Log")
        ws_audit.views.sheetView[0].showGridLines = True
        ws_audit.cell(row=1, column=1, value="VERSION HISTORY & REFINEMENT AUDIT LOG").font = title_font

        headers_audit = ["Version", "Timestamp (UTC)", "Patch Type", "Author / Agent", "Refinement Prompt / Details"]
        style_headers(ws_audit, headers_audit, start_row=3)

        patches = db.query(PlanPatch).filter(PlanPatch.transition_id == transition_id).order_by(PlanPatch.version_number.desc()).all()
        for idx, p in enumerate(patches, start=4):
            row_vals = [
                f"v{p.version_number}",
                p.applied_at.strftime("%Y-%m-%d %H:%M:%S"),
                p.patch_type,
                p.applied_by,
                p.nl_prompt or str(p.diff_payload),
            ]
            for c_idx, val in enumerate(row_vals, 1):
                cell = ws_audit.cell(row=idx, column=c_idx, value=val)
                cell.font = data_font
                cell.border = thin_border

        # Remove the blank initial sheet
        if default_sheet in wb.worksheets:
            wb.remove(default_sheet)

        # Autofit columns on all sheets
        for sheet in wb.worksheets:
            auto_fit_cols(sheet)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output


    @staticmethod
    def generate_schedule_csv(db: Session, transition_id: str) -> str:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "Session Title", "Level", "Scheduled Date", "Day", "Start Time",
            "End Time", "Duration Hours", "Delivery Mode", "Assigned SME", "Assigned Receiver",
            "Required Attendees", "Optional Attendees", "Status", "Conflicts"
        ])
        sessions = (
            db.query(KTSession)
            .filter(KTSession.transition_id == transition_id)
            .order_by(KTSession.scheduled_date, KTSession.start_time)
            .all()
        )
        stakeholders = db.query(Stakeholder).filter(Stakeholder.transition_id == transition_id).all()
        stk_map = {s.id: s for s in stakeholders}
        level_weights = {"L1": 1, "L2": 2, "L3": 3}

        for s in sessions:
            sme_obj = stk_map.get(s.sme_id)
            rcv_obj = stk_map.get(s.receiver_id)
            sme_name = sme_obj.name if sme_obj else "Unassigned"
            rcv_name = rcv_obj.name if rcv_obj else "Unassigned"

            t_rank = level_weights.get(s.level, 1)
            # Same level SME and Team member emails -> Required
            req_emails = [stk.email for stk in stakeholders if stk.email and level_weights.get(stk.level, 1) == t_rank]
            # Higher SME and Team member emails -> Optional
            opt_emails = [stk.email for stk in stakeholders if stk.email and level_weights.get(stk.level, 1) > t_rank]

            writer.writerow([
                s.session_title,
                s.level,
                s.scheduled_date.isoformat(),
                s.scheduled_date.strftime("%A"),
                s.start_time.strftime("%H:%M"),
                s.end_time.strftime("%H:%M"),
                s.duration_hours,
                s.delivery_mode,
                sme_name,
                rcv_name,
                "; ".join(req_emails),
                "; ".join(opt_emails),
                s.status,
                "; ".join(s.conflict_flags or []),
            ])
        return output.getvalue()

